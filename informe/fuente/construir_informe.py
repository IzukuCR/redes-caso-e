#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx"]
# ///
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(sys.argv[1])
SALIDA = Path(sys.argv[2])
AZUL = RGBColor(0x1F, 0x3A, 0x5F)
TOKEN = re.compile(r"(\*\*.+?\*\*|`.+?`|==.+?==)")


def leer(ruta):
    return (RAIZ / ruta).read_text(encoding="utf-8").rstrip("\n")


def sombrear(elemento_pr, color):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), color)
    elemento_pr.append(sh)


def fuente(run, nombre=None, tam=None, negrita=None, cursiva=None, color=None):
    if nombre:
        run.font.name = nombre
        rpr = run._element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), nombre)
    if tam:
        run.font.size = Pt(tam)
    if negrita is not None:
        run.bold = negrita
    if cursiva is not None:
        run.italic = cursiva
    if color is not None:
        run.font.color.rgb = color


def texto(p, cadena, tam=None, negrita=False):
    for parte in TOKEN.split(cadena):
        if not parte:
            continue
        if parte.startswith("**"):
            fuente(p.add_run(parte[2:-2]), tam=tam, negrita=True)
        elif parte.startswith("=="):
            r = p.add_run(parte[2:-2])
            fuente(r, tam=tam, negrita=True)
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif parte.startswith("`"):
            fuente(p.add_run(parte[1:-1]), nombre="Consolas", tam=(tam or 11) - 1)
        else:
            fuente(p.add_run(parte), tam=tam, negrita=negrita or None)


def campo(p, instruccion, resultado=""):
    for tipo in ("begin", None, "separate", "texto", "end"):
        r = p.add_run()
        if tipo == "texto":
            r.text = resultado
            continue
        if tipo is None:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = instruccion
            r._element.append(it)
            continue
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), tipo)
        r._element.append(fc)


def estilos(doc):
    st = doc.styles
    n = st["Normal"]
    n.font.name = "Calibri"
    n.font.size = Pt(11)
    n.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    n.paragraph_format.space_after = Pt(6)
    n.paragraph_format.line_spacing = 1.15
    for nombre, tam, antes in (("Heading 1", 16, 18), ("Heading 2", 13, 12), ("Heading 3", 11.5, 10)):
        h = st[nombre]
        h.font.name = "Calibri"
        h.font.size = Pt(tam)
        h.font.bold = True
        h.font.color.rgb = AZUL
        h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        h.paragraph_format.space_before = Pt(antes)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
    c = st.add_style("Codigo", 1)
    c.base_style = st["Normal"]
    c.font.name = "Consolas"
    c.font.size = Pt(8)
    c.element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    c.paragraph_format.space_after = Pt(0)
    c.paragraph_format.space_before = Pt(0)
    c.paragraph_format.line_spacing = 1.0
    c.paragraph_format.left_indent = Cm(0.2)
    sombrear(c.element.get_or_add_pPr(), "F1F3F6")
    cap = st["Caption"]
    cap.font.name = "Calibri"
    cap.font.size = Pt(9.5)
    cap.font.italic = True
    cap.font.bold = False
    cap.font.color.rgb = RGBColor(0x3D, 0x47, 0x56)
    cap.paragraph_format.space_after = Pt(8)


class Informe:
    def __init__(self):
        self.doc = Document()
        estilos(self.doc)
        s = self.doc.sections[0]
        s.page_width, s.page_height = Cm(21.59), Cm(27.94)
        s.left_margin = s.right_margin = Cm(2.4)
        s.top_margin = s.bottom_margin = Cm(2.3)
        s.different_first_page_header_footer = True
        self.tablas = 0
        self.figuras = 0
        self.titulos = []
        self.cabecera_pie(s)

    def cabecera_pie(self, s):
        hp = s.header.paragraphs[0]
        texto(hp, "Caso E: DMZ corporativa protegida | Grupo 5 | Redes Avanzadas I", tam=8.5)
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        texto(fp, "Página ", tam=9)
        campo(fp, "PAGE", "1")

    def h(self, nivel, cadena):
        self.titulos.append((nivel, cadena))
        return self.doc.add_heading(cadena, level=nivel)

    def p(self, cadena, alinear=None, tam=None, espacio=None):
        par = self.doc.add_paragraph()
        texto(par, cadena, tam=tam)
        if alinear == "centro":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if espacio is not None:
            par.paragraph_format.space_after = Pt(espacio)
        return par

    def salto(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def codigo(self, cadena):
        for linea in cadena.split("\n"):
            self.doc.add_paragraph(linea if linea else " ", style="Codigo")
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def leyenda(self, cadena, tipo):
        if tipo == "Tabla":
            self.tablas += 1
            n = self.tablas
        else:
            self.figuras += 1
            n = self.figuras
        par = self.doc.add_paragraph(style="Caption")
        texto(par, f"**{tipo} {n}.** {cadena}", tam=9.5)
        return par

    def tabla(self, titulo, cabecera, filas, anchos, tam=9):
        cap = self.leyenda(titulo, "Tabla")
        cap.paragraph_format.keep_with_next = True
        cap.paragraph_format.space_after = Pt(3)
        t = self.doc.add_table(rows=1, cols=len(cabecera))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for i, c in enumerate(cabecera):
            celda = t.rows[0].cells[i]
            celda.text = ""
            texto(celda.paragraphs[0], c, tam=tam, negrita=True)
            sombrear(celda._element.get_or_add_tcPr(), "D9E2F3")
        trpr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trpr.append(th)
        for fila in filas:
            cs = t.add_row().cells
            for i, c in enumerate(fila):
                cs[i].text = ""
                texto(cs[i].paragraphs[0], c, tam=tam)
        for r in t.rows:
            cant = OxmlElement("w:cantSplit")
            r._tr.get_or_add_trPr().append(cant)
            for i, ancho in enumerate(anchos):
                r.cells[i].width = Cm(ancho)
                for par in r.cells[i].paragraphs:
                    par.paragraph_format.space_after = Pt(1)
                    par.paragraph_format.space_before = Pt(1)
                    par.paragraph_format.line_spacing = 1.0
        total = sum(anchos)
        tbl_w = t._tbl.tblPr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            t._tbl.tblPr.append(tbl_w)
        tbl_w.set(qn("w:type"), "dxa")
        tbl_w.set(qn("w:w"), str(int(total / 2.54 * 1440)))
        for i, ancho in enumerate(anchos):
            t.columns[i].width = Cm(ancho)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def figura(self, ruta, ancho_cm, pie):
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(str(RAIZ / ruta), width=Cm(ancho_cm))
        c = self.leyenda(pie, "Figura")
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def marcador(self, id_, titulo, dispositivo, accion, debe):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = "Table Grid"
        celda = t.rows[0].cells[0]
        celda.width = Cm(14.5)
        sombrear(celda._element.get_or_add_tcPr(), "FFF2CC")
        celda.text = ""
        lineas = (
            f"==[{id_}]== Reemplazar este recuadro por la captura.",
            f"**Dispositivo:** {dispositivo}",
            f"**Comando o acción:** {accion}",
            f"**Debe verse:** {debe}",
        )
        for i, linea in enumerate(lineas):
            par = celda.paragraphs[0] if i == 0 else celda.add_paragraph()
            par.paragraph_format.space_before = Pt(3)
            par.paragraph_format.space_after = Pt(3)
            texto(par, linea, tam=9.5)
        self.figuras += 1
        c = self.doc.add_paragraph(style="Caption")
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        texto(c, f"**Figura {self.figuras}.** {titulo} Elaboración propia.", tam=9.5)

    def tabla_contenido(self, entradas):
        titulo = self.doc.add_paragraph()
        titulo.paragraph_format.space_after = Pt(10)
        fuente(titulo.add_run("Tabla de contenido"), tam=16, negrita=True, color=AZUL)
        parrafos = []
        for nivel, cad in entradas:
            par = self.doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.7 * (nivel - 1))
            par.paragraph_format.space_after = Pt(1)
            par.paragraph_format.line_spacing = 1.0
            texto(par, cad, tam=10, negrita=(nivel == 1))
            parrafos.append(par)
        if parrafos:
            primero = parrafos[0].runs[0]._r
            for t in ("begin", "instr", "separate"):
                r = OxmlElement("w:r")
                if t == "instr":
                    it = OxmlElement("w:instrText")
                    it.set(qn("xml:space"), "preserve")
                    it.text = 'TOC \\o "1-2" \\h \\z \\u'
                    r.append(it)
                else:
                    fc = OxmlElement("w:fldChar")
                    fc.set(qn("w:fldCharType"), t)
                    r.append(fc)
                primero.addprevious(r)
            r = OxmlElement("w:r")
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), "end")
            r.append(fc)
            parrafos[-1]._p.append(r)


def construir(inf):
    d = inf.doc

    portada(inf)
    inf.salto()
    inf.tabla_contenido(inf.indice_previo)
    inf.salto()

    inf.h(1, "1. Introducción y resumen ejecutivo")
    inf.p("Una empresa ficticia necesita publicar su servicio de nombres (DNS) y su sitio web (HTTP) hacia una red externa sin que esa apertura exponga su red interna. Este informe documenta el diseño, la implementación en GNS3 y la verificación de esa solución, asignada al grupo como Caso E, DMZ corporativa protegida, del curso Redes Avanzadas I (EIF 101O).")
    inf.p("La red se divide en tres zonas: una WAN que conecta con un proveedor de Internet simulado, una LAN interna con dos VLAN (usuarios y administración) y una DMZ donde vive el servidor SRV-DMZ, que ejecuta bind9 y nginx. Un FortiGate, llamado FGT-01, es el único punto de paso entre zonas: decide qué tráfico se permite, traduce direcciones en IPv4 y registra cada decisión. Hacia el exterior los equipos enrutan con rutas estáticas, y en el interior el FortiGate y el router R-LAN intercambian rutas con OSPFv2 y OSPFv3. Todo opera en IPv4 e IPv6, y en IPv6 el servidor se publica sin NAT.")
    inf.p("El diseño persigue cinco objetivos: publicar DNS y HTTP del servidor; bloquear cualquier acceso desde la red externa hacia la LAN; permitir la administración del servidor solo desde la VLAN de administración; contrastar una ACL de router con una política de firewall; y conservar evidencia de lo permitido y lo denegado.")
    inf.p("La maqueta tiene 11 nodos y ejecuta con el resultado esperado las 28 pruebas de la matriz del grupo, que incluye dos fallas controladas. La licencia de evaluación del FortiGate admite solo tres políticas de firewall, lo que obligó a rediseñar el conjunto de políticas (sección 6). La maqueta se construye con un script que parte de archivos de texto del repositorio, de modo que cualquier integrante puede reproducirla.")

    inf.h(1, "2. Requerimientos y decisiones de diseño")
    inf.h(2, "2.1 Requerimientos")
    inf.p("El enunciado del Caso E exige tres zonas lógicas como mínimo; la publicación de DNS y HTTP mediante políticas del FortiGate, con NAT o VIP en IPv4 y enrutamiento sin NAT66 en IPv6; el bloqueo de la red externa hacia la LAN; la administración del servidor desde una única zona autorizada; una ACL de router que se contraste con la política del firewall; y evidencia de logs o contadores. A esto se suman los requisitos comunes del curso: VLSM a partir de un bloque privado, direcciones globales y link-local de IPv6 con prefijos de documentación, un protocolo dinámico junto con una ruta estática, un FortiGate que controle tráfico real y siete pruebas mínimas de DNS, HTTP y política. El grupo adoptó el bloque sugerido para el quinto grupo: 10.50.0.0/16, 2001:db8:105::/48 y la zona g5.eif101o.test.")

    inf.h(2, "2.2 Arquitectura y justificación")
    inf.p("Tres principios guían el diseño. En IPv4, la LAN usa direcciones privadas (Rekhter et al., 1996) que nunca se anuncian hacia afuera: lo que sale pasa por NAT (Srisuresh y Egevang, 2001) y lo que entra solo alcanza la DMZ a través de una dirección virtual (VIP). En IPv6 no hay NAT66, de modo que el prefijo completo del grupo se enruta hasta el ISP y la única protección de la LAN es la política del firewall. Por último, la defensa es en profundidad: el firewall controla el tráfico entre zonas y las ACL de los routers cubren lo que el firewall no ve.")
    inf.p("El servidor se ubica en una DMZ porque es el único equipo accesible desde afuera. Si alguien lo compromete, la política que niega DMZ hacia LAN impide que salte a la red interna, y como no existe ninguna política de la DMZ hacia Internet, tampoco puede sacar datos (prueba E14). El FortiGate es el único equipo que toca las tres zonas, y los routers de cada lado cumplen una función propia: R-EDGE descarta en el borde el tráfico con origen falsificado antes de que consuma recursos del firewall, y R-LAN segmenta las VLAN, un tráfico que se enruta dentro del router y nunca llega al firewall.")
    inf.p("El servidor es un contenedor Docker con Alpine Linux y no una máquina virtual completa. El enunciado permite partir de una imagen Docker, el contenedor arranca en segundos y consume lo mínimo, y ejecuta los mismos servicios (bind9, nginx y OpenSSH) con la misma configuración que tendría en una VM. Como el curso pide no sumar equipos sin función, cada nodo de la tabla 1 responde a un requisito.")
    inf.tabla(
        "Inventario de equipos y función de cada uno.",
        ["Equipo", "Tipo", "Zona", "Función"],
        [
            ["R-ISP", "Router Cisco c7200", "Externa", "Simula Internet y enruta los prefijos públicos de la empresa"],
            ["R-EDGE", "Router Cisco c7200", "Borde", "Borde de la empresa con ACL anti-spoofing; solo se gestiona por consola"],
            ["FGT-01", "FortiGate VM", "WAN, LAN, DMZ", "Firewall de zonas: VIP, políticas, logs; OSPF con R-LAN"],
            ["R-LAN", "Router Cisco c7200", "LAN", "Router-on-a-stick para VLAN 10 y 20; ACL de segmentación y de gestión"],
            ["SW-LAN", "Switch Ethernet de GNS3", "LAN", "Troncal 802.1Q hacia R-LAN y puertos de acceso"],
            ["SW-DMZ", "Switch Ethernet de GNS3", "DMZ", "Segmento de la DMZ, permite sumar servidores sin tocar el firewall"],
            ["SRV-DMZ", "Contenedor Alpine", "DMZ", "DNS autoritativo (bind9), HTTP (nginx) y SSH: el servicio publicado"],
            ["PC-EXT", "Contenedor Alpine", "Externa", "Cliente de pruebas con dig, curl y nc"],
            ["PC-EXT-WEB", "Contenedor con navegador", "Externa", "Muestra la página por nombre desde afuera"],
            ["PC-USER", "Contenedor Debian con terminal y navegador", "LAN, VLAN 10", "Usuario interno: consume DNS y HTTP, prueba lo bloqueado"],
            ["PC-ADMIN", "Contenedor Debian con terminal y navegador", "LAN, VLAN 20", "Administra el servidor, R-LAN y el FortiGate; recibe el syslog"],
        ],
        [2.3, 3.7, 2.7, 8.1],
    )

    inf.h(2, "2.3 Plan de direccionamiento y VLSM")
    inf.p("El bloque 10.50.0.0/16 se reparte con VLSM (Fuller y Li, 2006): los segmentos se asignan de mayor a menor y de forma contigua, para no fragmentar el espacio libre. El tamaño de cada uno sale de un dimensionamiento supuesto: 200 usuarios necesitan una /24, que ofrece 254 direcciones útiles; 25 administradores caben en una /27, con 30; la DMZ, con hasta 10 servidores, en una /28, con 14; y el enlace entre el FortiGate y R-LAN, que solo tiene dos extremos, en una /30. Todo lo que está desde 10.50.1.52 en adelante queda libre para crecer.")
    inf.p("En IPv6, cada segmento recibe un /64 del /48 del grupo, como indica la arquitectura de direccionamiento (Hinden y Deering, 2006). El identificador de subred repite el número de VLAN (10 y 20) o de zona (30 para la DMZ), y los enlaces de tránsito usan ff00 y ff01 para reconocerse a simple vista. PC-USER obtiene su dirección por SLAAC (Thomson et al., 2007) a partir de los anuncios del router, mientras que el resto de los hosts usa direcciones fijas.")
    inf.tabla(
        "Segmentos internos con VLSM.",
        ["Segmento", "Hosts", "IPv4", "Máscara", "Rango útil", "IPv6"],
        [
            ["VLAN 10 Usuarios", "200", "10.50.0.0/24", "255.255.255.0", "10.50.0.1 a 10.50.0.254", "2001:db8:105:10::/64"],
            ["VLAN 20 Admin", "25", "10.50.1.0/27", "255.255.255.224", "10.50.1.1 a 10.50.1.30", "2001:db8:105:20::/64"],
            ["DMZ", "10", "10.50.1.32/28", "255.255.255.240", "10.50.1.33 a 10.50.1.46", "2001:db8:105:30::/64"],
            ["Tránsito FGT a R-LAN", "2", "10.50.1.48/30", "255.255.255.252", "10.50.1.49 a 10.50.1.50", "2001:db8:105:ff00::/64"],
            ["Reserva", "", "desde 10.50.1.52", "", "", "resto del /48"],
        ],
        [3.1, 1.2, 2.6, 2.7, 3.6, 3.6],
        tam=8.5,
    )
    inf.p("Los segmentos externos usan los rangos de documentación para IPv4 (Arkko et al., 2010) y el prefijo 2001:db8::/32 para IPv6 (Huston et al., 2004). La WAN pública de la empresa es 203.0.113.16/29 y allí vive el VIP 203.0.113.20; el enlace con el ISP es 203.0.113.0/30; y los clientes externos están en 198.51.100.0/24 (PC-EXT) y 192.0.2.0/24 (PC-EXT-WEB). El proveedor simulado usa 2001:db8:ffff::/48, fuera del /48 del grupo.")

    inf.h(2, "2.4 Implementación de DNS dual-stack y HTTP")
    inf.p("SRV-DMZ es servidor autoritativo de la zona g5.eif101o.test, con bind9 y la recursión desactivada: así no puede convertirse en un resolvedor abierto hacia Internet, y los clientes del laboratorio lo consultan directamente (Internet Systems Consortium, 2024). Escucha en 10.50.1.34 y en 2001:db8:105:30::10, por lo que atiende consultas por ambas familias, y nginx sirve la página en el puerto 80 sobre las mismas direcciones (nginx, s. f.).")
    inf.p("El servidor mantiene dos vistas DNS, una técnica conocida como split horizon. Desde afuera, `www` resuelve al VIP 203.0.113.20, que es la única dirección IPv4 que existe del lado de la WAN. Desde la LAN, ese VIP no es alcanzable sin una traducción adicional (hairpin NAT), así que la vista interna responde la IP real 10.50.1.34. El mecanismo funciona porque el VIP traduce el destino pero conserva el origen: el servidor ve la dirección real del cliente y elige la vista que le corresponde. La WAN pública, 2001:db8:105:ff01::/64, sale del mismo /48 que la LAN pero está fuera del firewall, por lo que se excluye de la vista interna. Los registros AAAA son idénticos en ambas vistas porque en IPv6 no hay traducción.")
    inf.p("Como `www` tiene registros A y AAAA, un navegador elige la familia por su cuenta y casi siempre prefiere IPv6 (Schinazi y Pauly, 2017). Para poder mostrar cada familia por nombre en la defensa se definieron `v4.g5.eif101o.test`, solo con registro A, y `v6.g5.eif101o.test`, solo con AAAA. La página imprime la IP del cliente mediante SSI de nginx, lo que demuestra el dual-stack y que el VIP no oculta el origen de la conexión.")

    inf.h(1, "3. Enrutamiento")
    inf.h(2, "3.1 Protocolo utilizado: OSPF")
    inf.p("OSPF es un protocolo de estado de enlace: cada router describe sus enlaces en anuncios (LSA), todos los routers de un área llegan a tener la misma base de datos y cada uno calcula con el algoritmo de Dijkstra el camino de menor costo hacia cada destino (Moy, 1998). OSPFv2 transporta rutas IPv4. OSPFv3 (Coltun et al., 2008) hace lo mismo para IPv6 con una diferencia que el curso destaca: las adyacencias se forman sobre las direcciones link-local, de modo que el siguiente salto que aprende cada router es la link-local del vecino y no su dirección global. Lo confirma la tabla de R-LAN de la sección 7, donde la ruta por defecto IPv6 sale por una dirección fe80.")
    inf.p("Se usa un área única, la 0, porque el foco del caso es la seguridad y con solo dos routers OSPF un área adicional no aportaría una función demostrable. Elegir OSPF en lugar de rutas estáticas internas tiene dos ventajas: una VLAN nueva en R-LAN se anuncia sola, y la caída del enlace entre FGT-01 y R-LAN es una falla controlada que se puede mostrar en vivo (prueba F02). El protocolo vive solo entre FGT-01 y R-LAN: del FortiGate hacia afuera todo es estático, como ocurre con un proveedor real, y R-EDGE y R-ISP nunca conocen 10.50.0.0/16, que solo existe detrás del NAT.")
    inf.tabla(
        "Parámetros OSPF (IPv4 e IPv6) de los dos routers internos.",
        ["Parámetro", "FGT-01", "R-LAN"],
        [
            ["Router-ID", "5.5.5.1", "5.5.5.2"],
            ["Área", "0", "0"],
            ["Interfaz activa (forma vecindad)", "port2", "Gi0/0 (Gi1/0 en el c7200)"],
            ["Interfaces pasivas (se anuncian sin vecinos)", "port3 (DMZ)", "subinterfaces de las VLAN 10 y 20"],
            ["Ruta por defecto", "La inyecta con default-information originate", "La aprende como externa tipo 2 (O*E2)"],
        ],
        [6.2, 5.3, 5.3],
    )

    inf.h(2, "3.2 Rutas estáticas")
    inf.p("Las rutas estáticas conectan lo que OSPF no toca. La tabla 4 las lista junto con su propósito. Las dos rutas blackhole del FortiGate merecen atención: sin la de 10.50.0.0/16, un destino interno inexistente seguiría la ruta por defecto y saldría a Internet; sin la de 2001:db8:105::/48, un /64 inexistente rebotaría entre el FortiGate y R-EDGE, que devuelve el /48 hacia el firewall. Las blackhole no tapan a las rutas de OSPF porque el prefijo más largo gana (10.50.0.0/24 prevalece sobre 10.50.0.0/16), y tampoco se redistribuyen.")
    inf.tabla(
        "Rutas estáticas de la maqueta.",
        ["Equipo", "Destino", "Siguiente salto", "Propósito"],
        [
            ["FGT-01", "0.0.0.0/0", "203.0.113.17", "Salida a Internet"],
            ["FGT-01", "::/0", "2001:db8:105:ff01::1", "Salida a Internet en IPv6"],
            ["FGT-01", "10.50.0.0/16", "blackhole, distancia 254", "Descartar destinos internos que no existen"],
            ["FGT-01", "2001:db8:105::/48", "blackhole, distancia 254", "Igual en IPv6"],
            ["R-EDGE", "0.0.0.0/0", "203.0.113.1", "Salida hacia el ISP"],
            ["R-EDGE", "::/0", "2001:db8:ffff:1::1", "Salida IPv6 hacia el ISP"],
            ["R-EDGE", "2001:db8:105::/48", "2001:db8:105:ff01::2", "Entregar el prefijo IPv6 de la empresa al FortiGate"],
            ["R-ISP", "203.0.113.16/29", "203.0.113.2", "Alcanzar el segmento público y el VIP"],
            ["R-ISP", "2001:db8:105::/48", "2001:db8:ffff:1::2", "Alcanzar el prefijo IPv6 de la empresa"],
        ],
        [2.1, 3.6, 4.4, 6.7],
    )
    inf.p("Una consulta externa recorre este camino: PC-EXT envía a R-ISP, que entrega a R-EDGE por la ruta estática hacia la /29; el FortiGate responde por el VIP, traduce el destino a 10.50.1.34 y entrega en la DMZ. En IPv6 el camino es el mismo, con la estática del /48 en R-EDGE y sin traducción. La respuesta vuelve por la sesión que el firewall ya tiene abierta.")

    inf.h(1, "4. Configuración de DNS y página web")
    inf.p("Los archivos completos de bind9 y de nginx están en el anexo B. Aquí se muestran los que definen el comportamiento descrito en la sección 2.4. La vista externa publica el VIP en IPv4 y la dirección global en IPv6.")
    inf.codigo(leer("servicios/dns/db.g5.eif101o.test.externa"))
    inf.p("La vista interna cambia solo los registros A, que apuntan a la IP real del servidor, y conserva los AAAA. El serial se sube en ambas vistas cada vez que cambia algo.")
    inf.codigo(leer("servicios/dns/db.g5.eif101o.test.interna"))
    inf.p("La página identifica al grupo, el caso y el FQDN, e imprime el nombre pedido, la IP del cliente y su familia, que sirve de evidencia en las pruebas P05 y P06. Se muestra el inicio del cuerpo del documento; el archivo completo, con sus estilos, está en el anexo B.")
    html = leer("servicios/web/index.html")
    cuerpo = html[html.index("<body") :].split("\n")[:26]
    inf.codigo("\n".join(cuerpo + ["  ..."]))

    inf.h(1, "5. Diagrama lógico y direccionamiento")
    inf.figura("informe/fuente/figuras/topologia.png", 14.3, "Diagrama lógico final de la maqueta del Caso E. Elaboración propia.")
    inf.p("Las interfaces de los routers siguen la nomenclatura del diseño (Gi0/0 a Gi0/2). En la imagen c7200 que usa la maqueta, la interfaz Gi0/N del diseño se llama Gi(N+1)/0: por ejemplo, Gi0/0 de R-LAN es Gi1/0 y sus subinterfaces son Gi2/0.10 y Gi2/0.20. La tabla 5 conserva la nomenclatura del diseño.")
    inf.tabla(
        "Direcciones por interfaz de los equipos de red.",
        ["Equipo", "Interfaz", "IPv4", "IPv6 global", "Link-local"],
        [
            ["R-ISP", "Gi0/0", "203.0.113.1/30", "2001:db8:ffff:1::1/64", "fe80::1"],
            ["R-ISP", "Gi0/1", "198.51.100.1/24", "2001:db8:ffff:100::1/64", "fe80::1"],
            ["R-ISP", "Gi0/2", "192.0.2.1/24", "2001:db8:ffff:200::1/64", "fe80::1"],
            ["R-EDGE", "Gi0/0", "203.0.113.2/30", "2001:db8:ffff:1::2/64", "fe80::2"],
            ["R-EDGE", "Gi0/1", "203.0.113.17/29", "2001:db8:105:ff01::1/64", "fe80::1"],
            ["FGT-01", "port1 (WAN)", "203.0.113.18/29", "2001:db8:105:ff01::2/64", "automática"],
            ["FGT-01", "VIP", "203.0.113.20 a 10.50.1.34", "no aplica", "no aplica"],
            ["FGT-01", "port2 (LAN)", "10.50.1.49/30", "2001:db8:105:ff00::1/64", "automática"],
            ["FGT-01", "port3 (DMZ)", "10.50.1.33/28", "2001:db8:105:30::1/64", "automática"],
            ["R-LAN", "Gi0/0", "10.50.1.50/30", "2001:db8:105:ff00::2/64", "fe80::2"],
            ["R-LAN", "Gi0/1.10", "10.50.0.1/24", "2001:db8:105:10::1/64", "fe80::1"],
            ["R-LAN", "Gi0/1.20", "10.50.1.1/27", "2001:db8:105:20::1/64", "fe80::1"],
        ],
        [1.8, 2.4, 4.2, 5.2, 3.2],
        tam=8.5,
    )
    inf.p("Los routers Cisco llevan fe80::1 fija en los gateways para que las tablas y capturas se lean con facilidad. Se repite en varios enlaces sin conflicto, porque una dirección link-local solo existe dentro de su enlace. El FortiGate usa link-local automáticas.")
    inf.tabla(
        "Direccionamiento de los hosts.",
        ["Host", "IPv4 y gateway", "IPv6 y gateway", "DNS configurado"],
        [
            ["SRV-DMZ", "10.50.1.34/28, gw 10.50.1.33", "2001:db8:105:30::10/64, gw 2001:db8:105:30::1", "él mismo"],
            ["PC-EXT", "198.51.100.10/24, gw .1", "2001:db8:ffff:100::10/64, gw fe80::1", "203.0.113.20 y 2001:db8:105:30::10"],
            ["PC-EXT-WEB", "192.0.2.10/24, gw .1", "2001:db8:ffff:200::10/64, gw fe80::1", "203.0.113.20 y 2001:db8:105:30::10"],
            ["PC-USER", "10.50.0.10/24, gw .1", "SLAAC sobre 2001:db8:105:10::/64, gw por RA", "10.50.1.34 y 2001:db8:105:30::10"],
            ["PC-ADMIN", "10.50.1.10/27, gw .1", "2001:db8:105:20::10/64, gw fe80::1", "10.50.1.34 y 2001:db8:105:30::10"],
        ],
        [2.2, 4.3, 5.8, 4.5],
        tam=8.5,
    )

    inf.h(1, "6. Diseño de seguridad")
    inf.h(2, "6.1 Zonas, NAT y políticas del FortiGate")
    inf.p("FGT-01 decide todo el tráfico entre WAN (port1), LAN (port2) y DMZ (port3). En IPv4 publica el servidor con un VIP que hace NAT estático 1:1 entre 203.0.113.20 y 10.50.1.34; en IPv6 lo publica por enrutamiento directo hacia 2001:db8:105:30::10, sin NAT66. El diseño original tenía seis políticas, pero la licencia de evaluación permanente de FortiOS limita las políticas de firewall (Fortinet, 2025): medido en el FGT-01, acepta tres y rechaza la cuarta con el aviso de que se alcanzó el máximo de entradas. Por eso la maqueta corre con un plan B de tres políticas, y el plan de seis queda en el repositorio como respaldo.")
    inf.p("El plan B se apoya en central NAT. FortiOS no deja mezclar un VIP con un objeto de dirección normal en el destino de una política, ni crear un VIP cuya IP externa sea igual a la mapeada. Con central NAT activado, el VIP solo traduce 203.0.113.20 hacia 10.50.1.34 y la política evalúa el tráfico ya traducido, de modo que un mismo objeto de destino atiende a los externos (por el VIP) y a los usuarios internos (directo a la IP real, como indica la vista interna del DNS). Así una sola política cubre WAN y LAN.")
    inf.tabla(
        "Matriz de políticas del FortiGate en vigor (plan B), en orden de evaluación.",
        ["ID", "Nombre", "Origen", "Destino", "Servicio", "Acción", "Justificación"],
        [
            ["1", "DNS-HTTP-A-DMZ", "port1 (all) y port2 (all)", "SRV-DMZ (10.50.1.34 y 2001:db8:105:30::10)", "DNS, HTTP", "Permitir, log de todas las sesiones", "Publicación controlada hacia afuera y consumo interno del servicio, sin gestión"],
            ["2", "ADMIN-A-DMZ", "port2, NET-ADMIN (VLAN 20)", "SRV-DMZ", "SSH, DNS, HTTP, PING", "Permitir, log", "Única zona autorizada para administrar el servidor"],
            ["3", "DENY-A-LAN", "port1 y port3 (all)", "all, hacia port2", "ALL", "Denegar, log de violaciones", "La LAN no queda expuesta desde WAN ni desde una DMZ comprometida"],
            ["0", "Implícita", "cualquiera", "all", "ALL", "Denegar, log activado", "Todo lo demás: SSH externo, SSH de usuarios, DMZ hacia Internet"],
        ],
        [0.9, 2.6, 2.7, 3.2, 1.9, 2.4, 3.1],
        tam=8,
    )
    inf.p("El diseño de seis políticas separaba lo que el plan B fusiona: PUB-DNS-HTTP (WAN a DMZ) y USUARIOS-A-DMZ (usuarios a DMZ) eran políticas distintas, DENY-EXT-A-LAN y DENY-DMZ-A-LAN eran dos denegaciones explícitas, y LAN-A-INTERNET permitía HTTP, HTTPS, DNS y ping con SNAT hacia port1. La salida de la LAN a Internet se eliminó en el plan B porque el caso no la exige. Con tres políticas, el requisito de mostrar qué política procesó cada tráfico se cumple igual: las denegaciones explícitas quedan en la política 3, las no publicadas en la implícita, y ambas están registradas.")
    inf.p("El acceso administrativo al FortiGate se limita por interfaz y por cuenta. En port1 solo se permite ping, sin HTTPS ni SSH; en port2, HTTPS, SSH y ping; en port3, solo ping. Además, la cuenta admin tiene trusted hosts restringidos a 10.50.1.0/27 y 2001:db8:105:20::/64, así que aunque PC-USER alcance port2 el equipo ignora su intento de gestión. Esto funciona solo si todas las cuentas de administrador tienen trusted hosts, por lo que también se revisó la cuenta creada por defecto. Los intentos de gestión denegados no pasan por una política de reenvío, de modo que se activó el log de denegaciones local-in para que queden registrados.")

    inf.h(2, "6.2 ACL de router")
    inf.p("Hay tres ACL, cada una en IPv4 e IPv6, y todas cubren huecos que el firewall no puede ver. Las líneas de denegación llevan `log` para que el router deje su propio registro además de los contadores.")
    inf.tabla(
        "ACL-SEG-USUARIOS en R-LAN, entrada de la subinterfaz de la VLAN 10.",
        ["Línea", "Acción", "Origen a destino", "Protocolo", "Motivo"],
        [
            ["10", "Permitir", "Usuarios a Admin", "ICMP echo-reply", "Que Admin pueda hacer ping a Usuarios y reciba la respuesta"],
            ["20", "Permitir", "Usuarios a Admin", "TCP established", "Respuestas a conexiones que inicia Admin"],
            ["30", "Denegar, log", "Usuarios a Admin", "IP", "Usuarios no inician nada hacia Administración"],
            ["40", "Permitir", "any a any", "IP", "El resto sigue hacia el FortiGate"],
        ],
        [1.3, 2.2, 3.4, 3.0, 6.9],
        tam=8.5,
    )
    inf.p("Las líneas 10 y 20 existen porque la ACL no tiene estado: sin ellas, Admin no recibiría las respuestas. La línea 20 tiene una limitación conocida, pues `established` solo revisa que el segmento traiga el indicador ACK o RST, de modo que un usuario que forje un segmento con ACK lo hace llegar a la VLAN 20. El host lo descarta porque no tiene una conexión abierta, pero la ACL no lo frena. Se dejó así a propósito: ilustra por qué el control fino entre zonas corresponde al firewall con estado.")
    inf.tabla(
        "ACL-VTY-ADMIN en R-LAN y ACL-BORDE-IN en R-EDGE.",
        ["ACL", "Línea", "Acción", "Coincidencia", "Motivo"],
        [
            ["VTY-ADMIN (líneas vty)", "10", "Permitir", "NET-ADMIN, IPv4 e IPv6", "Única zona que administra el router"],
            ["VTY-ADMIN (líneas vty)", "20", "Denegar, log", "any", "Todo otro intento de SSH queda registrado"],
            ["BORDE-IN (Gi0/0, desde el ISP)", "10 a 13", "Denegar, log", "Origen 10/8, 172.16/12, 192.168/16, 127/8", "Anti-spoofing: rangos privados no vienen de Internet (Ferguson y Senie, 2000)"],
            ["BORDE-IN", "20", "Denegar, log", "Origen 203.0.113.16/29", "Nadie de afuera usa el prefijo propio"],
            ["BORDE-IN", "30 a 33", "Denegar, log", "TCP 22 y 23 hacia 203.0.113.2 y .17", "Nadie gestiona el router de borde desde Internet"],
            ["BORDE-IN", "40", "Permitir", "any a any", "El FortiGate decide el resto"],
        ],
        [3.6, 1.6, 2.2, 4.4, 5.0],
        tam=8.5,
    )
    inf.p("Las líneas vty aceptan solo SSH, nunca Telnet. R-EDGE se gestiona únicamente por consola. En IPv6 la ACL de borde deniega como origen el propio 2001:db8:105::/48 y fc00::/7.")
    inf.h(2, "6.3 ACL frente a política del firewall")
    inf.p("La diferencia central es el estado. Una ACL evalúa cada paquete por separado, así que la respuesta necesita su propia regla y un paquete forjado puede pasar si cumple el patrón, como el segmento con ACK de la línea 20. Una política del FortiGate pertenece a una sesión: la respuesta se acepta porque la solicitud fue permitida, y un paquete que no pertenece a ninguna sesión se descarta. La ACL se aplica a una interfaz y un sentido, y filtra por IP, protocolo y puerto; la política se aplica a un par de interfaces, usa objetos y servicios, traduce direcciones y registra por sesión el ID de la política que decidió. Por eso en este diseño las ACL cubren el anti-spoofing, la segmentación entre VLAN y la gestión de los routers, y el firewall controla el paso entre zonas y la publicación del servicio.")
    inf.h(2, "6.4 Conservación de evidencia")
    inf.p("El FortiGate guarda sus logs en el disco local de 30 GB de la plantilla y los copia por syslog (UDP 514) a PC-ADMIN, de modo que sobreviven a un reinicio del equipo; R-LAN hace lo mismo con el log de sus ACL. R-EDGE conserva solo el buffer local porque enviar syslog a la LAN obligaría a abrir un acceso WAN a LAN, justo lo que el caso prohíbe. El syslog no va a SRV-DMZ porque es el servidor expuesto y un atacante podría borrar la evidencia. Los contadores de cada línea de ACL y los aciertos de cada política completan el registro.")

    inf.h(1, "7. Evidencias de verificación")
    inf.p("Las pruebas se ejecutaron en vivo sobre la maqueta con los scripts `pruebas/verificar_maqueta.py` y `pruebas/verificar_extras.py`, que guardan la salida completa en `pruebas/evidencias/` (maqueta-20261024-1734.txt y extras-20261024-1813.txt). Las 28 pruebas de la matriz dieron el resultado esperado. Los fragmentos de este capítulo son copia literal de esa salida.")
    inf.h(2, "7.1 Enrutamiento")
    inf.p("La prueba R01 confirma las dos adyacencias OSPF en estado Full entre FGT-01 y R-LAN. En OSPFv3 el vecino se identifica por el router-ID, y el siguiente salto de las rutas es una link-local.")
    inf.codigo(
        "FGT-01 # get router info ospf neighbor\n"
        "Neighbor ID     Pri   State           Dead Time   Address         Interface\n"
        "5.5.5.2           1   Full/Backup     00:00:33    10.50.1.50      port2\n\n"
        "FGT-01 # get router info6 ospf neighbor\n"
        "Neighbor ID     Pri   State           Dead Time   Interface  Instance ID\n"
        "5.5.5.2           1   Full/Backup     00:00:30    port2      0"
    )
    inf.p("La prueba R02 compara las tablas de los cinco equipos con el diseño y no faltó ninguna ruta. Como muestra, estas son las rutas que R-LAN aprende por OSPF: la ruta por defecto externa (O*E2) y la DMZ, en IPv4 e IPv6.")
    inf.codigo(
        "R-LAN#show ip route\n"
        "O*E2  0.0.0.0/0 [110/10] via 10.50.1.49, 00:35:12, GigabitEthernet1/0\n"
        "O        10.50.1.32/28 [110/2] via 10.50.1.49, 00:35:12, GigabitEthernet1/0\n\n"
        "R-LAN#show ipv6 route\n"
        "OE2 ::/0 [110/10]\n"
        "     via FE80::EC5:6BFF:FEAE:1, GigabitEthernet1/0\n"
        "O   2001:DB8:105:30::/64 [110/2]\n"
        "     via FE80::EC5:6BFF:FEAE:1, GigabitEthernet1/0"
    )
    inf.p("En FGT-01, la tabla confirma la ruta por defecto hacia R-EDGE, las redes de la LAN aprendidas por OSPF desde R-LAN, las redes conectadas y la ruta resumen hacia Null que descarta los destinos internos inexistentes (E15). La salida completa, del 27 de octubre de 2026, está en el Anexo G.")
    inf.h(2, "7.2 DNS por IPv4 e IPv6")
    inf.p("Las cuatro primeras pruebas mínimas se ejecutan desde PC-EXT contra el servidor publicado. P01 y P02 resuelven los registros A y AAAA; P03 y P04 verifican que el transporte DNS responde por IPv4 (al VIP) y por IPv6 (a la dirección global), con la respuesta autoritativa de la vista externa.")
    inf.codigo(
        "PC-EXT$ dig +short @203.0.113.20 www.g5.eif101o.test A\n203.0.113.20\n"
        "PC-EXT$ dig +short @203.0.113.20 www.g5.eif101o.test AAAA\n2001:db8:105:30::10\n"
        "PC-EXT$ dig -4 +short @203.0.113.20 www.g5.eif101o.test A\n203.0.113.20\n"
        "PC-EXT$ dig -6 +short @2001:db8:105:30::10 www.g5.eif101o.test A\n203.0.113.20"
    )
    inf.p("Desde la LAN (E06), PC-USER consulta al servidor y recibe la vista interna: `dig +short @10.50.1.34 www.g5.eif101o.test A` devuelve 10.50.1.34.")
    inf.h(2, "7.3 HTTP por nombre")
    inf.p("P05 y P06 piden la página por nombre sobre cada familia y obtienen HTTP 200. Las capturas del navegador de PC-EXT-WEB, con la IP del cliente que imprime la página, acompañan la prueba.")
    inf.codigo(
        "PC-EXT$ curl -4 -s -m 6 -o /dev/null -w 'http4=%{http_code}\\n' http://www.g5.eif101o.test/\nhttp4=200\n"
        "PC-EXT$ curl -6 -s -m 6 -o /dev/null -w 'http6=%{http_code}\\n' http://www.g5.eif101o.test/\nhttp6=200"
    )
    inf.marcador("CAP-01", "Página del grupo abierta por IPv4.", "PC-EXT-WEB (navegador del nodo, interfaz gráfica de GNS3)", "Abrir http://v4.g5.eif101o.test/ y dejar visible la barra de direcciones.", "La página del grupo con la IP del cliente en IPv4, el nombre del servidor y los nombres de los integrantes.")
    inf.marcador("CAP-02", "Página del grupo abierta por IPv6.", "PC-EXT-WEB (navegador del nodo, interfaz gráfica de GNS3)", "Abrir http://v6.g5.eif101o.test/ y dejar visible la barra de direcciones.", "La misma página con la IP del cliente en IPv6 (prefijo 2001:db8:ffff:...).")
    inf.h(2, "7.4 Políticas de seguridad: qué política procesó el tráfico")
    inf.p("Los logs de Forward Traffic identifican la política que decidió cada sesión (P07). Los dos primeros fragmentos muestran tráfico permitido por la política 1: una sesión HTTP externa, con el destino traducido por el VIP (`trandisp=dnat`), y una sesión HTTP interna hacia la IP real, sin traducción. El tercero muestra una denegación registrada en la política implícita (`policyid=0`), tomada cuando la política 1 estaba desactivada en la prueba F01.")
    inf.codigo(
        "srcip=198.51.100.10 srcintf=\"port1\" dstip=203.0.113.20 dstport=80 dstintf=\"port3\"\n"
        "  action=\"close\" policyid=1 policyname=\"DNS-HTTP-A-DMZ\" service=\"HTTP\" trandisp=\"dnat\" tranip=10.50.1.34\n\n"
        "srcip=10.50.0.10 srcintf=\"port2\" dstip=10.50.1.34 dstport=80 dstintf=\"port3\"\n"
        "  action=\"close\" policyid=1 policyname=\"DNS-HTTP-A-DMZ\" service=\"HTTP\" trandisp=\"noop\"\n\n"
        "srcip=198.51.100.10 srcintf=\"port1\" dstip=203.0.113.20 dstport=80 dstintf=\"port3\"\n"
        "  action=\"deny\" policyid=0 service=\"HTTP\" trandisp=\"dnat\" tranip=10.50.1.34"
    )
    inf.p("Las políticas 2 y 3 también quedan identificadas en el log. La política 2 procesó el SSH de PC-ADMIN al servidor (E04); la política 3, DENY-A-LAN, negó el ping externo a la LAN por IPv6 (E01) y el ping del servidor a PC-USER (E09). Una sesión TCP que el cliente cierra con RST aparece como `client-rst` y no es un fallo.")
    inf.codigo("""srcip=10.50.1.10 srcintf="port2" dstip=10.50.1.34 dstport=22 dstintf="port3"
  proto=6 action="client-rst" policyid=2

srcip=2001:db8:ffff:100::10 srcintf="port1" dstip=2001:db8:105:10::1 dstintf="port2"
  proto=58 action="deny" policyid=3 policyname="DENY-A-LAN" service="PING6"

srcip=10.50.1.34 srcintf="port3" dstip=10.50.0.10 dstintf="port2"
  proto=1 action="deny" policyid=3 policyname="DENY-A-LAN" service="PING" """.strip())
    inf.p("Los contadores de la tabla de políticas (`diagnose firewall iprope show 100004 N` para IPv4 e `iprope6` para IPv6) son la segunda evidencia que pide el enunciado: cada política acumula aciertos con la fecha del primero y del último. Son los de la sesión de pruebas del 27 de octubre de 2026; la salida de la política 3 está en el Anexo G.")
    inf.tabla(
        "Aciertos acumulados por política en FGT-01.",
        ["Política", "Función", "Aciertos IPv4", "Aciertos IPv6", "Prueba que los explica"],
        [
            ["1", "DNS-HTTP-A-DMZ (permitir)", "157", "158", "P01 a P06 y E06"],
            ["2", "Administración SSH al servidor (permitir)", "1", "0", "E04"],
            ["3", "DENY-A-LAN (denegar)", "2", "2", "E09 por IPv4; E01 por IPv6"],
            ["0", "Denegación implícita", "3", "0", "E03 (SSH externo)"],
        ],
        [1.6, 5.6, 2.4, 2.4, 4.8],
        tam=9,
    )
    inf.marcador("CAP-03", "Forward Traffic del FortiGate por política.", "PC-ADMIN, Firefox en la consola gráfica (VNC) del nodo, hacia la GUI de FGT-01 (https://10.50.1.49, IP de port2)", "Log & Report, Forward Traffic. Aplicar el filtro Policy ID = 1, luego 2, luego 3 y luego 0 (una captura o un montaje con las cuatro). Si no hay filas, repetir antes E01, E03, E04 y E09 de la matriz.", "Columnas Source, Destination, Policy ID, Policy Name y Action. Debe verse Accept para las políticas 1 y 2, y Deny para la 3 (DENY-A-LAN) y la 0.")
    inf.h(2, "7.5 Pruebas específicas del caso E")
    inf.tabla(
        "Resultado observado en las pruebas de bloqueo y de acceso restringido.",
        ["ID", "Prueba", "Resultado observado"],
        [
            ["E01", "Externo hacia la LAN por IPv6 (ping a 2001:db8:105:10::1)", "100% de pérdida; registrado en la política 3"],
            ["E02", "Externo hacia la LAN por IPv4 (10.50.0.10)", "Destination Host Unreachable desde R-ISP: no hay ruta"],
            ["E02b", "Igual, con rutas 10.50.0.0/16 temporales en R-ISP y R-EDGE", "El paquete llega al FortiGate y se deniega en port1; las rutas se retiran después"],
            ["E03", "SSH externo al VIP, puerto 22", "Operation timed out; cae en la denegación implícita"],
            ["E04", "PC-ADMIN por SSH al servidor", "Puerto 22 abierto (política 2)"],
            ["E05", "PC-USER por SSH al servidor", "Operation timed out; denegación implícita"],
            ["E07", "PC-USER hacia PC-ADMIN (ping)", "Packet filtered desde 10.50.0.1; línea 30 de ACL-SEG-USUARIOS"],
            ["E08", "PC-ADMIN hacia PC-USER (ping)", "Responde, 64 bytes, ttl 63, gracias a la línea 10"],
            ["E09", "SRV-DMZ hacia PC-USER", "100% de pérdida; política 3"],
            ["E10", "Origen falsificado 10.99.99.99 desde R-ISP", "0 de 3 respuestas; contador de la línea 10 de ACL-BORDE-IN en 3"],
            ["E11", "HTTPS al FortiGate por port1 desde PC-EXT", "Sin respuesta (https=000)"],
            ["E12", "SSH a R-LAN desde PC-USER y desde PC-ADMIN", "Bloqueado para PC-USER; abierto para PC-ADMIN"],
            ["E13", "GUI del FortiGate desde PC-USER y desde PC-ADMIN", "gui=000 para PC-USER (trusted hosts); gui=200 para PC-ADMIN"],
            ["E14", "SRV-DMZ hacia 198.51.100.10", "100% de pérdida; denegación implícita"],
            ["E15", "PC-USER hacia 10.50.99.9 y 2001:db8:105:99::9", "Descartado por la blackhole del FortiGate, sin salir a Internet"],
            ["E16", "Logs fuera del equipo", "Syslog de PC-ADMIN recibió la denegación de E01, la de E07 y la de E12"],
        ],
        [1.2, 7.0, 8.6],
        tam=8.5,
    )
    inf.p("Los contadores y registros de las ACL confirman las pruebas E10 y E16:")
    inf.codigo(
        "R-EDGE#show ip access-lists ACL-BORDE-IN\n"
        "    10 deny ip 10.0.0.0 0.255.255.255 any log (3 matches)\n"
        "    40 permit ip any any (12 matches)\n\n"
        "syslog de PC-ADMIN:\n"
        "<190>102: *Oct 25 00:11:49.055: %SEC-6-IPACCESSLOGDP: list ACL-SEG-USUARIOS denied icmp 10.50.0.10 -> 10.50.1.10 (8/0), 1 packet\n"
        "<190>103: *Oct 25 00:12:49.055: %SEC-6-IPACCESSLOGNP: list ACL-VTY-ADMIN denied 0 10.50.0.10 -> 0.0.0.0, 1 packet"
    )
    inf.marcador("CAP-04", "GUI del FortiGate abierta desde PC-ADMIN.", "PC-ADMIN, Firefox en la consola gráfica (VNC) del nodo; y la consola auxiliar (terminal) de PC-USER para el segundo comando", "En PC-ADMIN abrir https://10.50.1.49 e iniciar sesión. En PC-USER correr: curl -k -s -m 5 -o /dev/null -w gui=%{http_code} https://10.50.1.49", "En PC-ADMIN, el Dashboard del FortiGate (gui=200). En PC-USER, la salida gui=000, que muestra el bloqueo por trusted hosts.")
    inf.h(2, "7.6 Fallas y cambios controlados")
    inf.p("**F01, desactivar la política 1.** Con la política desactivada, la publicación externa cae (`http4=000`) y el log muestra la denegación implícita; al reactivarla, vuelve a funcionar (`http4=200`) sin otro cambio. La prueba muestra que la publicación depende de esa única política y que la denegación por defecto protege el servicio mientras no esté permitido.")
    inf.p("**F02, apagar la interfaz de R-LAN hacia el FortiGate.** La vecindad OSPF cae, al instante en R-LAN y al vencer el dead interval en el FortiGate, y la LAN pierde la ruta por defecto y la DMZ: el HTTP de los usuarios pasa a `http4=000`. Desde afuera la publicación sigue funcionando (`http4=200`), porque no depende de R-LAN. Al levantar la interfaz, la vecindad vuelve a Full y el HTTP de los usuarios regresa a `http4=200`.")

    inf.h(1, "8. Troubleshooting")
    inf.p("Durante el desarrollo se registraron las siguientes incidencias, en orden de aparición.")
    inf.h(3, "Incidencia 1: el FortiGate no reenviaba tráfico")
    inf.p("Con la imagen de FortiGate 8.0.0, ningún paquete atravesaba el equipo aunque respondía ping a sus propias interfaces. El diagnóstico empezó por descartar recursos: el motor IPS mostraba cero procesos de trabajo, un crashlog hablaba de que los recursos de la VM excedían el límite de la licencia con 2048 MB, y ajustar la RAM no resolvió nada. Una prueba aislada con solo dos interfaces y una política que permitía todo falló igual en 8.0.0 y en 7.6.7, y `get system status` mostraba License Status Invalid y un límite de 1 CPU. La documentación de Fortinet explica la causa: una VM recién desplegada ya no trae una licencia de evaluación válida y debe activarse la evaluación permanente con una cuenta de FortiCloud (Fortinet, 2025). La solución fue automatizar ese paso en el script de la maqueta (subcomando `licenciar`), que conecta un nodo NAT temporal a un puerto del FortiGate, ejecuta `execute vm-license-options` y `execute vm-license`, espera el reinicio y verifica License Status Valid. Un fallo posterior del script, que informaba que el FortiGate no tenía salida a Internet por el puerto de licencia, resultó ser de tiempo y no de topología: una espera fija de 20 segundos era corta para el DHCP del nodo NAT recién iniciado, y se cambió por un sondeo con límite de 150 segundos.")
    inf.h(3, "Incidencia 2: la licencia solo admite tres políticas")
    inf.p("Al cargar el diseño de seis políticas, el FortiGate rechazó la cuarta con el mensaje de que se había alcanzado el máximo de entradas. Se midió el límite: acepta exactamente tres políticas, mientras que las rutas estáticas (cuatro) y las tres interfaces se aplicaron sin problema. Se rediseñó el conjunto fusionando políticas (sección 6.1). Al fusionarlas apareció una segunda restricción: FortiOS rechaza un VIP mezclado con un objeto normal en el destino de una política (\"entry not found in datasource\") y también un VIP con IP externa igual a la mapeada. La salida fue activar central NAT, de modo que el VIP solo traduzca y la política use la IP real. El cambio se verificó con las pruebas P01 a P07 y E01 a E16 sobre el plan B.")
    inf.h(3, "Incidencia 3: IPv4 desaparecía en PC-USER")
    inf.p("Al configurar PC-USER, que entonces usaba una imagen Alpine, con una dirección fija IPv4 y un bloque `inet6 auto` para SLAAC, el host quedaba sin IPv4. El `ifup` de BusyBox, que usa Alpine, no soporta `inet6 auto` y aborta todo el archivo de interfaces. Se quitó el bloque inet6: la dirección IPv6 llega igual por SLAAC porque el kernel acepta los anuncios del router (`accept_ra`). Al pasar después PC-USER a una imagen Debian se conservó la configuración y SLAAC siguió funcionando.")
    inf.h(3, "Incidencia 4: disco de logs corrupto en el FortiGate")
    inf.p("El disco de logs del FortiGate resultó estar corrupto: la partición ext4 del disco de la plantilla tenía el sistema de archivos dañado. Se resolvió con `execute disk format 16`, que reinicia el equipo; la licencia siguió válida y la configuración persistió. La verificación posterior mostró 329 logs en el disco, entre ellos los de las pruebas.")
    inf.h(3, "Incidencia 5: nombres de interfaz distintos a los del diseño")
    inf.p("La imagen de router disponible es un c7200, y no el IOSv que suponía el diseño, y sus interfaces no se llaman Gi0/N sino Gi(N+1)/0, de modo que los nombres del diseño no coincidían con los del equipo al cargar las configuraciones. Se documentó la equivalencia (Gi0/0 es Gi1/0; las subinterfaces de R-LAN son Gi2/0.10 y Gi2/0.20). Los archivos de `configs/` conservan los nombres Gi0/N del diseño y el script `gns3ctl.py` los traduce al cargarlos; la cabecera de cada archivo lo explica, y quien pegue las configuraciones a mano debe usar el nombre real del equipo. La matriz anota el nombre real.")

    inf.h(3, "Incidencia 6: PC-ADMIN no tenía navegador")
    inf.p("El diseño pedía en PC-ADMIN un terminal y un navegador para abrir la GUI del FortiGate, pero la plantilla Ubuntu Docker Guest que se había elegido solo trae terminal. Como todas las pruebas automáticas usan curl y ssh, nadie lo notó hasta intentar abrir la GUI. Se construyó la imagen `caso-e/webterm`, que parte de gns3/webterm (Debian con Firefox por VNC) y le suma dnsutils y netcat para que las pruebas sigan corriendo; el terminal queda en la consola auxiliar del nodo. PC-ADMIN y PC-USER usan ahora esa imagen, con la misma red y las mismas direcciones, y la prueba E13 (GUI desde PC-ADMIN sí, desde PC-USER no) se repitió con el resultado anterior.")

    inf.h(1, "9. Conclusiones, recomendaciones y mejoras futuras")
    inf.h(2, "9.1 Conclusiones técnicas")
    inf.p("El diseño cumple los requisitos del Caso E: el servidor publica DNS y HTTP por IPv4 y por IPv6 con nombre, la red externa no alcanza la LAN por ninguna de las dos familias, solo la VLAN de administración llega al servidor por SSH, y cada decisión del FortiGate queda registrada con su política. Las 28 pruebas de la matriz dieron el resultado esperado, incluidas dos fallas controladas que muestran cómo se comporta la red cuando se apaga una política y cuando cae una adyacencia OSPF.")
    inf.p("Tres lecciones destacan. La primera es que la política de un firewall y una ACL no son intercambiables: la ACL cubrió el anti-spoofing y la segmentación, pero su límite con `established` mostró por qué el control con estado pertenece al firewall. La segunda, que en IPv6 sin NAT la protección de la LAN depende por completo de la política, y por eso las pruebas E01 y E15 importan. La tercera, que las restricciones de la licencia son parte del diseño: la cuarta política no cabía, y repensar el conjunto con central NAT produjo un diseño más compacto sin perder ninguna prueba.")
    inf.p("Para el curso, el proyecto exigió integrar equipos de dos fabricantes, razonar el direccionamiento desde el requisito y no desde la tabla, y documentar de forma que otro integrante pueda reproducir el trabajo, en vez de depender de la máquina de una persona.")
    inf.h(2, "9.2 Recomendaciones y limitaciones")
    inf.p("Con la licencia de evaluación, la maqueta ejecuta el plan B de tres políticas y la LAN no tiene salida a Internet (la política de salida del plan original se eliminó). Con una licencia completa conviene volver al diseño de seis políticas, que separa por zonas y deja explícita cada denegación. Otras limitaciones conocidas: el DNS de PC-USER se configura a mano y solo por IPv4 (los registros AAAA se resuelven igual), la línea `established` de la ACL de usuarios admite segmentos con ACK forjado, y R-EDGE conserva su log solo en el buffer local.")
    inf.h(2, "9.3 Mejoras futuras")
    inf.p("La mejora de mayor valor es sustituir la ACL de segmentación por un firewall por zonas (ZBF) o una ACL reflexiva en R-LAN, para que el control entre VLAN tenga estado. Convendría también entregar el DNS de los clientes IPv6 por RDNSS o DHCPv6 sin estado, incorporar perfiles de inspección del FortiGate (IPS y filtrado web) sobre la política de publicación, centralizar los logs con marca de tiempo sincronizada por NTP, y añadir un segundo servidor en la DMZ con alta disponibilidad para el servicio publicado.")

    inf.h(1, "10. Referencias")
    refs = [
        "Arkko, J., Cotton, M., y Vegoda, L. (2010). *IPv4 address blocks reserved for documentation* (RFC 5737). Internet Engineering Task Force. https://www.rfc-editor.org/info/rfc5737",
        "Coltun, R., Ferguson, D., Moy, J., y Lindem, A. (2008). *OSPF for IPv6* (RFC 5340). Internet Engineering Task Force. https://www.rfc-editor.org/info/rfc5340",
        "Ferguson, P., y Senie, D. (2000). *Network ingress filtering: Defeating denial of service attacks which employ IP source address spoofing* (RFC 2827). Internet Engineering Task Force. https://www.rfc-editor.org/info/rfc2827",
        "Fortinet. (2025). *Permanent trial mode for FortiGate-VM*. En *FortiOS 7.4.6 administration guide*. Fortinet Document Library. https://docs.fortinet.com/document/fortigate/7.4.6/administration-guide/441460",
        "Fuller, V., y Li, T. (2006). *Classless inter-domain routing (CIDR): The Internet address assignment and aggregation plan* (RFC 4632). Internet Engineering Task Force. https://www.rfc-editor.org/info/rfc4632",
        "Hinden, R., y Deering, S. (2006). *IP version 6 addressing architecture* (RFC 4291). Internet Engineering Task Force. https://www.rfc-editor.org/info/rfc4291",
        "Huston, G., Lord, A., y Smith, P. (2004). *IPv6 address prefix reserved for documentation* (RFC 3849). Internet Engineering Task Force. https://www.rfc-editor.org/info/rfc3849",
        "Internet Systems Consortium. (2024). *BIND 9 administrator reference manual*. https://bind9.readthedocs.io",
        "Moy, J. (1998). *OSPF version 2* (RFC 2328). Internet Engineering Task Force. https://www.rfc-editor.org/info/rfc2328",
        "nginx. (s. f.). *nginx documentation*. https://nginx.org/en/docs/",
        "Rekhter, Y., Moskowitz, B., Karrenberg, D., de Groot, G. J., y Lear, E. (1996). *Address allocation for private internets* (RFC 1918). Internet Engineering Task Force. https://www.rfc-editor.org/info/rfc1918",
        "Schinazi, D., y Pauly, T. (2017). *Happy Eyeballs version 2: Better connectivity using concurrency* (RFC 8305). Internet Engineering Task Force. https://www.rfc-editor.org/info/rfc8305",
        "Srisuresh, P., y Egevang, K. (2001). *Traditional IP network address translator (Traditional NAT)* (RFC 3022). Internet Engineering Task Force. https://www.rfc-editor.org/info/rfc3022",
        "Thomson, S., Narten, T., y Jinmei, T. (2007). *IPv6 stateless address autoconfiguration* (RFC 4862). Internet Engineering Task Force. https://www.rfc-editor.org/info/rfc4862",
    ]
    for r in refs:
        par = d.add_paragraph()
        par.paragraph_format.left_indent = Cm(1.0)
        par.paragraph_format.first_line_indent = Cm(-1.0)
        par.paragraph_format.space_after = Pt(5)
        partes = re.split(r"(\*.+?\*)", r)
        for x in partes:
            if x.startswith("*") and x.endswith("*") and len(x) > 2:
                fuente(par.add_run(x[1:-1]), tam=10, cursiva=True)
            else:
                fuente(par.add_run(x), tam=10)

    anexos(inf)


def anexos(inf):
    inf.salto()
    inf.h(1, "Anexo A. Configuraciones de los equipos")
    inf.p("Copia literal de los archivos del repositorio en `configs/`. Las credenciales son las del laboratorio (admin y admin). El FortiGate se muestra en su versión en vigor, el plan B; el plan A de seis políticas está en `configs/FGT-01.conf`.")
    for titulo, ruta in (
        ("A.1 R-ISP", "configs/R-ISP.txt"),
        ("A.2 R-EDGE", "configs/R-EDGE.txt"),
        ("A.3 R-LAN", "configs/R-LAN.txt"),
        ("A.4 FGT-01, plan B", "configs/FGT-01-planB.conf"),
    ):
        inf.h(2, titulo)
        inf.codigo(leer(ruta))
    inf.h(1, "Anexo B. Servicios de SRV-DMZ")
    inf.h(2, "B.1 named.conf de bind9")
    inf.codigo(leer("servicios/dns/named.conf"))
    inf.h(2, "B.2 Configuración de nginx")
    inf.codigo(leer("servicios/web/caso-e.conf"))
    inf.h(2, "B.3 index.html completo")
    inf.codigo(leer("servicios/web/index.html"))
    inf.h(1, "Anexo C. Switches")
    inf.p("Los switches de GNS3 no tienen línea de comandos. SW-LAN lleva el puerto 0 como troncal 802.1Q (VLAN nativa 1, VLAN 10 y 20 etiquetadas), el puerto 1 en acceso para la VLAN 10 (PC-USER) y el puerto 2 en acceso para la VLAN 20 (PC-ADMIN). SW-DMZ lleva el puerto 0 hacia port3 del FortiGate y el puerto 1 hacia SRV-DMZ, ambos en acceso por la VLAN 1.")
    inf.h(1, "Anexo D. Comandos de verificación")
    inf.codigo(
        "# Desde PC-EXT\n"
        "dig +short @203.0.113.20 www.g5.eif101o.test A\n"
        "dig +short @203.0.113.20 www.g5.eif101o.test AAAA\n"
        "dig -4 +short @203.0.113.20 www.g5.eif101o.test A\n"
        "dig -6 +short @2001:db8:105:30::10 www.g5.eif101o.test A\n"
        "curl -4 http://www.g5.eif101o.test/\n"
        "curl -6 http://www.g5.eif101o.test/\n"
        "nc -zv -w3 203.0.113.20 22\n\n"
        "# FortiGate\n"
        "get router info ospf neighbor\n"
        "get router info6 ospf neighbor\n"
        "get router info routing-table all\n"
        "get router info6 routing-table\n"
        "execute log filter device disk\n"
        "execute log filter category 0\n"
        "execute log display\n"
        "diagnose firewall iprope list 100004\n\n"
        "# Routers Cisco\n"
        "show ip route\n"
        "show ipv6 route\n"
        "show ip access-lists ACL-BORDE-IN\n"
        "show logging | include IPACCESSLOG"
    )
    inf.h(1, "Anexo E. Proyecto GNS3")
    inf.p("Enlace de descarga del proyecto exportado sin imágenes propietarias (Archivo, Exportar proyecto): https://drive.google.com/drive/folders/1mmUOyVEecsfnWPkWwIrvBDGNry5oO2gS?usp=sharing.")
    inf.p("La maqueta también puede reconstruirse desde cero con el script `automatizacion/gns3ctl.py` del repositorio, cuya guía de uso está en `automatizacion/README.md`.")
    inf.h(1, "Anexo F. Matriz de pruebas")
    inf.p("Las 28 pruebas con el resultado esperado y el observado. La evidencia completa está en `pruebas/evidencias/`.")
    inf.tabla(
        "Matriz de pruebas del Caso E.",
        ["ID", "Prueba", "Esperado", "Observado"],
        [
            ['P01', 'Registro A desde PC-EXT', '203.0.113.20', '203.0.113.20'],
            ['P02', 'Registro AAAA desde PC-EXT', '2001:db8:105:30::10', '2001:db8:105:30::10'],
            ['P03', 'Consulta DNS por IPv4', 'Respuesta autoritativa; política 1', 'Respuesta de la vista externa; accept en la política 1'],
            ['P04', 'Consulta DNS por IPv6', 'Respuesta autoritativa; política 1', 'Respuesta de la vista externa; accept en la política 1'],
            ['P05', 'HTTP por IPv4 con nombre', 'HTTP 200, IP de cliente IPv4', 'http4=200'],
            ['P06', 'HTTP por IPv6 con nombre', 'HTTP 200, IP de cliente IPv6', 'http6=200'],
            ['P07', 'Política permitida y denegada', 'P05 permitida y E01 denegada', 'Política 1 acepta; política 3 deniega'],
            ['E01', 'Externo hacia LAN por IPv6', 'Bloqueado', '100% de pérdida; política 3'],
            ['E02', 'Externo hacia LAN por IPv4', 'Sin ruta', 'Destination Host Unreachable desde R-ISP'],
            ['E02b', 'Igual, con ruta temporal', 'Bloqueado en el FortiGate', 'Se deniega en port1'],
            ['E03', 'SSH externo al servidor', 'Denegación implícita', 'Timeout; política 0'],
            ['E04', 'PC-ADMIN por SSH al servidor', 'Conecta', 'Puerto 22 abierto; política 2'],
            ['E05', 'PC-USER por SSH al servidor', 'Denegación implícita', 'Timeout'],
            ['E06', 'PC-USER usa DNS y HTTP', 'Vista interna', '10.50.1.34; HTTP 200; política 1'],
            ['E07', 'PC-USER hacia PC-ADMIN', 'ACL-SEG-USUARIOS, línea 30', 'Packet filtered desde 10.50.0.1'],
            ['E08', 'PC-ADMIN hacia PC-USER', 'Responde por la línea 10', 'Responde, ttl 63'],
            ['E09', 'SRV-DMZ hacia PC-USER', 'Bloqueado', '100% de pérdida; política 3'],
            ['E10', 'Origen falsificado en el borde', 'ACL-BORDE-IN, línea 10', '0 de 3 respuestas; contador en 3'],
            ['E11', 'HTTPS al FortiGate por port1', 'Rechazado', 'https=000'],
            ['E12', 'SSH a R-LAN', 'Bloqueado a PC-USER; abierto a PC-ADMIN', 'Bloqueado y abierto, respectivamente'],
            ['E13', 'GUI del FortiGate', 'Solo PC-ADMIN', 'gui=000 en PC-USER; gui=200 en PC-ADMIN'],
            ['E14', 'SRV-DMZ hacia Internet', 'Denegación implícita', '100% de pérdida'],
            ['E15', 'Destino interno inexistente', 'Descartado por blackhole', 'Descartado sin salir a Internet'],
            ['E16', 'Logs fuera del equipo', 'Syslog en PC-ADMIN', 'Recibió las denegaciones de E01, E07 y E12'],
            ['R01', 'Vecinos OSPFv2 y OSPFv3', 'Full en ambos', 'Full/Backup en port2'],
            ['R02', 'Tablas de rutas', 'Coinciden con el diseño', 'No faltó ninguna ruta'],
            ['F01', 'Desactivar la política 1', 'La publicación cae y vuelve', 'http4=000 y luego http4=200'],
            ['F02', 'Apagar la interfaz de R-LAN hacia el FortiGate', 'La LAN pierde servicio; el externo sigue', 'http4=000 en la LAN; http4=200 desde afuera'],
        ],
        [1.2, 4.8, 4.8, 6.0],
        tam=8.5,
    )

    inf.h(1, "Anexo G. Salidas del FortiGate")
    inf.h(2, "G.1 Tabla de rutas de FGT-01")
    inf.codigo("""FGT-01 # get router info routing-table all
S*      0.0.0.0/0 [10/0] via 203.0.113.17, port1
S       10.50.0.0/16 [254/0] is a summary, Null
O       10.50.0.0/24 [110/2] via 10.50.1.50, port2
O       10.50.1.0/27 [110/2] via 10.50.1.50, port2
C       10.50.1.32/28 is directly connected, port3
C       10.50.1.48/30 is directly connected, port2
C       203.0.113.16/29 is directly connected, port1

FGT-01 # get router info6 routing-table
S*      ::/0 [10/0] via 2001:db8:105:ff01::1, port1
S       2001:db8:105::/48 [254/0] via ::, Null
O       2001:db8:105:10::/64 [110/2] via fe80::2, port2
O       2001:db8:105:20::/64 [110/2] via fe80::2, port2
C       2001:db8:105:30::/64 via ::, port3
C       2001:db8:105:ff00::/64 via ::, port2
C       2001:db8:105:ff01::/64 via ::, port1""")
    inf.h(2, "G.2 Contadores de la política 3")
    inf.codigo("""FGT-01 # diagnose firewall iprope show 100004 3
idx:3
pkts:2 (2 0 0 0 0 0 0 0)
hit count:2 (2 0 0 0 0 0 0 0)
    first hit:2026-10-27 13:21:36 last hit:2026-10-27 13:21:37

FGT-01 # diagnose firewall iprope6 show 100004 3
idx:3
pkts:2 (2 0 0 0 0 0 0 0)
hit count:2 (2 0 0 0 0 0 0 0)
    first hit:2026-10-27 13:21:27 last hit:2026-10-27 13:21:28""")

def portada(inf):
    d = inf.doc
    for texto_, tam, neg in (
        ("Universidad Nacional", 16, True),
        ("Facultad de Ciencias Exactas y Naturales", 12, False),
        ("Escuela de Informática", 12, False),
    ):
        par = inf.p(texto_, alinear="centro", tam=tam, espacio=2)
        par.runs[0].bold = neg
    for _ in range(4):
        d.add_paragraph()
    par = inf.p("Redes Avanzadas I, EIF 101O", alinear="centro", tam=13)
    par = inf.p("Proyecto del curso, Parte I", alinear="centro", tam=13, espacio=18)
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fuente(t.add_run("Caso E: DMZ corporativa protegida"), tam=26, negrita=True, color=AZUL)
    par = inf.p("Diseño, implementación y verificación en GNS3 de una red empresarial con publicación controlada de DNS y HTTP", alinear="centro", tam=13, espacio=30)
    inf.p("**Grupo 5**", alinear="centro", tam=13, espacio=10)
    for nombre in ("Joanfer Hidalgo Chaves", "Roy Arias Mejia", "Isaac Rodriguez Aguero", "Eduardo Canton Grajales"):
        inf.p(nombre, alinear="centro", tam=12, espacio=2)
    for _ in range(3):
        d.add_paragraph()
    inf.p("Profesor: Mag. Pablo Sandoval Barrantes", alinear="centro", tam=12, espacio=2)
    inf.p("31 de octubre de 2026", alinear="centro", tam=12, espacio=14)
    inf.p("Proyecto GNS3 descargable: https://drive.google.com/drive/folders/1mmUOyVEecsfnWPkWwIrvBDGNry5oO2gS?usp=sharing", alinear="centro", tam=11)


def activar_actualizacion(doc):
    settings = doc.settings.element
    nodo = OxmlElement("w:updateFields")
    nodo.set(qn("w:val"), "true")
    settings.append(nodo)


def main():
    previo = Informe()
    previo.indice_previo = []
    construir(previo)
    entradas = [
        (n, t)
        for n, t in previo.titulos
        if n <= 2 and not (n == 2 and t[:2] in ("A.", "B.", "C.", "D.", "E.", "F.", "G."))
    ]
    final = Informe()
    final.indice_previo = entradas
    construir(final)
    final.doc.core_properties.title = "Caso E: DMZ corporativa protegida"
    final.doc.core_properties.author = "Grupo 5"
    activar_actualizacion(final.doc)
    final.doc.save(str(SALIDA))
    print(SALIDA)


if __name__ == "__main__":
    main()
